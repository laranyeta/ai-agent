import openpyxl
from mcp.server.mcpserver import MCPServer

mcp = MCPServer("job-offers")

@mcp.tool()
def get_offers() -> str:
    try:
        wb = openpyxl.load_workbook("ofertas.xlsx")
        ws = wb.active
        
        offers = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0]:
                role, company, description = row
                offers.append(f"**{role}** in {company}\n{description}")
        
        if not offers:
            return "No saved offers have been found."
        
        return "\n\n---\n\n".join(offers)
    except FileNotFoundError:
        return "[ERROR] .xlsx file has not been found."

@mcp.tool()
def get_keyword(keyword: str) -> str:
    try:
        wb = openpyxl.load_workbook("jobs.xlsx")
        ws = wb.active
        
        found = []
        word = keyword.lower()
        
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0]:
                role, company, description = row
                if word in role.lower() or word in descripcion.lower():
                    found.append(f"**{role}** in {company}\n{description}")
        
        if not found:
            return f"No offers with keyword "{keyword}" have been found."
        
        return "\n\n---\n\n".join(found)
    except FileNotFoundError:
        return "[ERROR] .xlsx file has not been found."

if __name__ == "__main__":
    mcp.run(transport="stdio")
